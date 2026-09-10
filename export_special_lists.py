# -*- coding: utf-8 -*-
import sys
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
"""
从 records.json 中导出"空号 / 停机"专项 Excel 清单
供电话催收团队优先处理无法接通的用户。
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, 'out')

TARGET_TAGS = {'空号', '停机'}


def main():
    records_path = os.path.join(OUT, 'records.json')
    if not os.path.exists(records_path):
        print("[WARN] 未找到 out/records.json，请先完成同步。")
        return

    try:
        data = json.load(open(records_path, 'r', encoding='utf-8'))
    except Exception as e:
        print(f"[FAIL] 读取 records.json 失败: {repr(e)}")
        import traceback
        traceback.print_exc()
        raise

    # 用协议号 aid 关联用户信息（主要是手机号）
    user_by_aid = {}
    for r in data.get('rows', []):
        user_by_aid[r.get('id')] = {
            'phone': r.get('ph', ''),
            'uid': r.get('uid', ''),
            'level': r.get('lv', ''),
            'lname': r.get('ln', ''),
            'city': r.get('ci', ''),
            'area': r.get('ar', ''),
        }

    rows = []
    for idx, rec in enumerate(data.get('reception_detail', [])):
        raw_tags = rec.get('tags') or []
        if not isinstance(raw_tags, (list, tuple, set)):
            print(f"[WARN] 第 {idx} 条接待记录 tags 格式异常，跳过：{raw_tags!r}")
            continue
        tags = set(raw_tags)
        hit = tags & TARGET_TAGS
        if not hit:
            continue
        info = user_by_aid.get(rec.get('aid'), {})
        rows.append({
            '协议号': rec.get('aid') or '',
            '用户ID': (rec.get('uid') or info.get('uid') or ''),
            '手机号': info.get('phone', ''),
            '标签': '、'.join(sorted(hit)),
            '接待时间': rec.get('time', ''),
            '接待类型': rec.get('type', ''),
            '接待人': rec.get('solver', ''),
            '接待内容': rec.get('detail', ''),
            '用户等级': info.get('level', ''),
            '用户名称': info.get('lname', ''),
            '城市': info.get('city', ''),
            '区域': info.get('area', ''),
        })

    if not rows:
        print("ℹ 本次同步未找到空号/停机标签记录，未生成专项清单。")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "空号停机明细"

    headers = ['协议号', '用户ID', '手机号', '标签', '接待时间', '接待类型',
               '接待人', '接待内容', '用户等级', '用户名称', '城市', '区域']
    ws.append(headers)

    for r in rows:
        ws.append([r[h] for h in headers])

    # 表头样式
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 自动列宽（限制最大 60）
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

    out_path = os.path.join(OUT, '特殊清单_空号停机.xlsx')
    try:
        wb.save(out_path)
        print(f"[OK] 已生成空号/停机专项清单：{out_path}  （共 {len(rows)} 条记录）")
    except Exception as e:
        print(f"[FAIL] 保存 Excel 失败: {repr(e)}")
        import traceback
        traceback.print_exc()
        raise


if __name__ == '__main__':
    try:
        main()
    except Exception:
        import traceback
        traceback.print_exc()
        raise
